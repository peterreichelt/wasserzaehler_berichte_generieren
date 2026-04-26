from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from watermeter.config import AppConfig, ExcelConfig, MappingOverrides, PlanRadarConfig
from watermeter.exceptions import PlanRadarApiError
from watermeter.importing.models import DiscoverySummary, MeterReading
from watermeter.importing.pipeline import run_pipeline


# Erzeugt eine wiederverwendbare Testkonfiguration für die Pipeline.
def _build_config(tmp_path: Path, *, bulk: bool) -> AppConfig:
    return AppConfig(
        input_excel_path=tmp_path / "input.xlsx",
        output_dir=tmp_path / "output",
        customer_id="1531527",
        target_project_id=None if bulk else "wddpbya",
        target_project_name=None,
        target_meter_location=None if bulk else "HAUPT_Z_1",
        log_level="INFO",
        excel=ExcelConfig(
            sheet_name=None,
            columns={
                "meter_number": "A",
                "reading_date": "B",
                "reading_value": "C",
                "property_name": "I",
                "meter_location": "N",
            },
        ),
        planradar=PlanRadarConfig(
            include_attachments=True,
            attachments_enabled=True,
            attachment_download_timeout_seconds=60.0,
            request_pause_seconds=2.1,
            mapping_overrides=MappingOverrides(),
        ),
    )


# Schreibt eine kleine historische Excel-Datei für die Pipeline-Tests.
def _write_excel(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["W 1", date(2026, 1, 10), 111, None, None, None, None, None, "Property One", None, None, None, None, "HAUPT_Z_1"])
    sheet.append(["W 2", date(2026, 1, 11), 222, None, None, None, None, None, "Property Two", None, None, None, None, "HAUPT_Z_2"])
    workbook.save(path)


# Baut eine einfache Discovery-Zusammenfassung für Stub-Ergebnisse.
def _build_summary(project_id: str, project_name: str) -> DiscoverySummary:
    return DiscoverySummary(
        project_id=project_id,
        project_name=project_name,
        ticket_type_id="tt-1",
        ticket_type_name="Wasserzählerablesung neu",
        field_ids={},
        project_list_ids=[],
        notes=[],
    )


def test_run_pipeline_bulk_mode_collects_excel_matches_and_remaining_projects(tmp_path: Path, monkeypatch) -> None:
    config = _build_config(tmp_path, bulk=True)
    _write_excel(config.input_excel_path)

    class DummyClient:
        pass

    project_one = {"id": "p1", "attributes": {"name": "Property One"}}
    project_three = {"id": "p3", "attributes": {"name": "Property Three"}}

    def fake_from_env_files(**kwargs):  # noqa: ANN003
        return DummyClient()

    def fake_load_projects(client, customer_id):  # noqa: ANN001
        return [project_one, project_three]

    def fake_load_project_readings(client, cfg, project, *, allowed_meter_locations=None):  # noqa: ANN001
        if project["id"] == "p1":
            assert allowed_meter_locations == {"HAUPT_Z_1"}
            readings = [
                MeterReading(
                    source="planradar",
                    property_name="Property One",
                    meter_location="HAUPT_Z_1",
                    reading_date=date(2026, 2, 1),
                    reading_value=333,
                    project_id="p1",
                    ticket_id="t1",
                )
            ]
            return type("Result", (), {"summary": _build_summary("p1", "Property One"), "readings": readings})()
        assert allowed_meter_locations is None
        readings = [
            MeterReading(
                source="planradar",
                property_name="Property Three",
                meter_location="HAUPT_Z_9",
                reading_date=date(2026, 2, 2),
                reading_value=444,
                project_id="p3",
                ticket_id="t3",
            )
        ]
        return type("Result", (), {"summary": _build_summary("p3", "Property Three"), "readings": readings})()

    downloaded: list[list[MeterReading]] = []
    downloaded_covers: list[tuple[str, str]] = []

    def fake_download(client, readings, cfg):  # noqa: ANN001
        downloaded.append(list(readings))

    def fake_download_cover(client, cfg, *, project_id, property_name):  # noqa: ANN001
        downloaded_covers.append((project_id, property_name))

    monkeypatch.setattr("watermeter.importing.pipeline.PlanRadarClient.from_env_files", fake_from_env_files)
    monkeypatch.setattr("watermeter.importing.pipeline.load_projects", fake_load_projects)
    monkeypatch.setattr("watermeter.importing.pipeline.load_project_readings", fake_load_project_readings)
    monkeypatch.setattr("watermeter.importing.pipeline.download_attachments_for_readings", fake_download)
    monkeypatch.setattr("watermeter.importing.pipeline.download_project_cover_image", fake_download_cover)

    output_path, summary = run_pipeline(config)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    rows = list(workbook["combined_readings"].iter_rows(values_only=True))
    assert output_path.name == "all_properties__combined_readings.xlsx"
    assert workbook.sheetnames == ["combined_readings"]
    assert len(rows) == 5
    assert {rows[1][1], rows[2][1], rows[3][1], rows[4][1]} == {
        "Property One",
        "Property Two",
        "Property Three",
    }
    assert summary.project_id == "bulk"
    assert len(downloaded) == 1
    assert {reading.property_name for reading in downloaded[0]} == {"Property One", "Property Three"}
    assert set(downloaded_covers) == {("p1", "Property One"), ("p3", "Property Three")}


def test_run_pipeline_focused_mode_keeps_discovery_sheet(tmp_path: Path, monkeypatch) -> None:
    config = _build_config(tmp_path, bulk=False)
    _write_excel(config.input_excel_path)

    class DummyClient:
        pass

    downloaded_covers: list[tuple[str, str]] = []

    def fake_from_env_files(**kwargs):  # noqa: ANN003
        return DummyClient()

    def fake_discover(client, cfg):  # noqa: ANN001
        readings = [
            MeterReading(
                source="planradar",
                property_name="Property One",
                meter_location="HAUPT_Z_1",
                reading_date=date(2026, 2, 1),
                reading_value=333,
                project_id="p1",
                ticket_id="t1",
            )
        ]
        return type("Result", (), {"summary": _build_summary("p1", "Property One"), "readings": readings})()

    monkeypatch.setattr("watermeter.importing.pipeline.PlanRadarClient.from_env_files", fake_from_env_files)
    monkeypatch.setattr("watermeter.importing.pipeline.discover_and_load_readings", fake_discover)
    monkeypatch.setattr("watermeter.importing.pipeline.download_attachments_for_readings", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        "watermeter.importing.pipeline.download_project_cover_image",
        lambda client, cfg, *, project_id, property_name: downloaded_covers.append((project_id, property_name)),
    )

    output_path, _summary = run_pipeline(config)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert "discovery" in workbook.sheetnames
    assert downloaded_covers == [("p1", "Property One")]


def test_run_pipeline_bulk_mode_skips_projects_without_matching_ticket_type(
    tmp_path: Path,
    monkeypatch,
) -> None:
    config = _build_config(tmp_path, bulk=True)
    _write_excel(config.input_excel_path)

    class DummyClient:
        pass

    project_one = {"id": "p1", "attributes": {"name": "Property One"}}
    project_bad = {"id": "p2", "attributes": {"name": "Property Bad"}}

    def fake_from_env_files(**kwargs):  # noqa: ANN003
        return DummyClient()

    def fake_load_projects(client, customer_id):  # noqa: ANN001
        return [project_one, project_bad]

    def fake_load_project_readings(client, cfg, project, *, allowed_meter_locations=None):  # noqa: ANN001
        if project["id"] == "p2":
            raise PlanRadarApiError("Could not discover a suitable ticket type for water meter readings.")
        readings = [
            MeterReading(
                source="planradar",
                property_name="Property One",
                meter_location="HAUPT_Z_1",
                reading_date=date(2026, 2, 1),
                reading_value=333,
                project_id="p1",
                ticket_id="t1",
            )
        ]
        return type("Result", (), {"summary": _build_summary("p1", "Property One"), "readings": readings})()

    monkeypatch.setattr("watermeter.importing.pipeline.PlanRadarClient.from_env_files", fake_from_env_files)
    monkeypatch.setattr("watermeter.importing.pipeline.load_projects", fake_load_projects)
    monkeypatch.setattr("watermeter.importing.pipeline.load_project_readings", fake_load_project_readings)
    monkeypatch.setattr("watermeter.importing.pipeline.download_attachments_for_readings", lambda *args, **kwargs: None)
    downloaded_covers: list[tuple[str, str]] = []
    monkeypatch.setattr(
        "watermeter.importing.pipeline.download_project_cover_image",
        lambda client, cfg, *, project_id, property_name: downloaded_covers.append((project_id, property_name)),
    )

    output_path, summary = run_pipeline(config)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    rows = list(workbook["combined_readings"].iter_rows(values_only=True))
    assert summary.project_id == "bulk"
    assert output_path.name == "all_properties__combined_readings.xlsx"
    assert {rows[1][1], rows[2][1], rows[3][1]} == {"Property One", "Property Two"}
    assert downloaded_covers == [("p1", "Property One")]
