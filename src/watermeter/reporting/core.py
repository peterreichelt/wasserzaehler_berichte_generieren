from __future__ import annotations

import base64
import logging
import mimetypes
import os
import tempfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from importlib.resources import files
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageOps

from ..config import AppConfig, ReportingConfig
from ..exceptions import ExcelDataError, ReportingError
from ..shared.normalization import ensure_directory, normalize_text, sanitize_filename


os.environ.setdefault("MPLCONFIGDIR", os.path.join(tempfile.gettempdir(), "wasserzaehler-mpl"))
import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt


logger = logging.getLogger(__name__)


TABLE_ROWS_SINGLE_PAGE_ONE_COL = 12
TABLE_ROWS_PER_COL = 14
TABLE_COLS_PER_PAGE = 2
CHART_ROWS_PER_PAGE = 2
COMBINED_PAGE_MAX_TABLE_ROWS = 8
CHART_WIDTH_IN = 11.8
CHART_HEIGHT_PER_ROW_IN = 2.35
CHART_DPI = 180
CHART_Y_PADDING_RATIO = 0.10


@dataclass
class ReportReading:
    """Beschreibt eine einzelne Ablesung im Reporting-Kontext."""

    source: str
    property_name: str
    meter_location: str
    meter_number: str | None
    reading_date: date
    reading_value: float
    attachment_paths: list[Path] = field(default_factory=list)
    notes: str | None = None
    days_since_previous: int | None = None
    daily_consumption: float | None = None
    meter_change: bool = False


@dataclass
class PhotoEntry:
    """Beschreibt ein einzelnes Foto innerhalb des Reports."""

    image_path: Path
    photo_number: int
    meter_number: str | None
    reading_date: date


@dataclass
class ReadingBand:
    """Bündelt ein Diagrammband inklusive linker und rechter Tabellenhälfte."""

    band_number: int
    readings: list[ReportReading]
    left_readings: list[ReportReading]
    right_readings: list[ReportReading]


@dataclass
class MeterReport:
    """Beschreibt einen vollständigen Bericht für Liegenschaft und Zählerplatz."""

    property_name: str
    meter_location: str
    meter_numbers: list[str]
    readings: list[ReportReading]
    bands: list[ReadingBand]
    photo_pages: list[list[PhotoEntry]]
    cover_image_path: Path | None
    generated_on: date
    version_code: str

    @property
    def date_range(self) -> str:
        """Liefert den Datumsbereich des Reports als Text."""
        if not self.readings:
            return "-"
        return f"{format_date(self.readings[0].reading_date)} - {format_date(self.readings[-1].reading_date)}"

    @property
    def summary_meter_number(self) -> str:
        """Liefert die Zählernummern für kompakte Zusammenfassungen."""
        if not self.meter_numbers:
            return self.meter_location
        return ", ".join(self.meter_numbers)


def format_date(value: date | None) -> str:
    """Formatiert ein Datum im deutschen Format."""
    if value is None:
        return ""
    return value.strftime("%d.%m.%Y")


def format_number(value: float | int | None, decimals: int = 0) -> str:
    """Formatiert Zahlen mit deutschem Dezimaltrennzeichen."""
    if value is None:
        return ""
    number = float(value)
    if decimals == 0 and number.is_integer():
        return f"{int(number)}"
    return f"{number:.{decimals}f}".replace(".", ",")


def format_consumption(value: float | None) -> str:
    """Formatiert einen Verbrauchswert mit zwei Nachkommastellen."""
    if value is None:
        return ""
    return f"{float(value):.2f}".replace(".", ",")


def split_list(items: list[Any], chunk_size: int) -> list[list[Any]]:
    """Teilt eine Liste in gleich große Teilmengen auf."""
    if chunk_size <= 0:
        raise ValueError("chunk_size muss größer als 0 sein.")
    return [items[index : index + chunk_size] for index in range(0, len(items), chunk_size)]


def image_file_to_data_uri(path: Path) -> str:
    """Lädt eine Datei als Data-URI."""
    if not path.exists():
        return ""
    mime_type, _ = mimetypes.guess_type(str(path))
    if not mime_type:
        mime_type = "image/png"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def raster_image_file_to_data_uri(
    path: Path,
    *,
    max_width: int,
    max_height: int,
    jpeg_quality: int = 88,
) -> str:
    """Optimiert eine Rastergrafik und kodiert sie als Data-URI."""
    if not path.exists():
        return ""
    mime_type, _ = mimetypes.guess_type(str(path))
    if mime_type == "image/svg+xml":
        return image_file_to_data_uri(path)

    with Image.open(path) as image:
        prepared = ImageOps.exif_transpose(image)
        prepared.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)
        has_alpha = "A" in prepared.getbands()

        buffer = BytesIO()
        if has_alpha:
            prepared.save(buffer, format="PNG", optimize=True)
            encoded_mime_type = "image/png"
        else:
            prepared.convert("RGB").save(
                buffer,
                format="JPEG",
                optimize=True,
                quality=jpeg_quality,
                subsampling=0,
            )
            encoded_mime_type = "image/jpeg"

    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:{encoded_mime_type};base64,{encoded}"


def _pil_image_to_data_uri(image: Image.Image, mime_type: str = "image/png") -> str:
    """Kodiert ein Pillow-Bild als Data-URI."""
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def load_combined_readings_workbook(path: Path) -> list[ReportReading]:
    """Lädt die kombinierte Excel-Arbeitsmappe für das Reporting."""
    if not path.exists():
        raise ReportingError(f"Die kombinierte Arbeitsmappe existiert nicht: {path}")
    logger.info("Kombinierte Arbeitsmappe wird aus %s geladen", path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "combined_readings" not in workbook.sheetnames:
        raise ReportingError("Die Arbeitsmappe enthält kein Blatt 'combined_readings'.")

    sheet = workbook["combined_readings"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ReportingError("Das Blatt 'combined_readings' ist leer.")

    header = [normalize_text(item) for item in rows[0]]
    index_map = {name: idx for idx, name in enumerate(header)}
    required = [
        "source",
        "property_name",
        "meter_location",
        "meter_number",
        "reading_date",
        "reading_value",
        "attachment_paths",
    ]
    missing = [name for name in required if name not in index_map]
    if missing:
        raise ReportingError(f"Der Arbeitsmappe fehlen erforderliche Spalte(n): {', '.join(missing)}")

    base_dir = path.parent
    readings: list[ReportReading] = []
    for row in rows[1:]:
        if not any(row):
            continue
        reading_date = _coerce_report_date(row[index_map["reading_date"]])
        raw_attachments = normalize_text(row[index_map["attachment_paths"]])
        attachment_paths = [
            (base_dir / part.strip()).resolve()
            for part in raw_attachments.split(";")
            if part and part.strip()
        ]
        readings.append(
            ReportReading(
                source=normalize_text(row[index_map["source"]]),
                property_name=normalize_text(row[index_map["property_name"]]),
                meter_location=normalize_text(row[index_map["meter_location"]]),
                meter_number=normalize_text(row[index_map["meter_number"]]) or None,
                reading_date=reading_date,
                reading_value=float(row[index_map["reading_value"]]),
                attachment_paths=attachment_paths,
                notes=normalize_text(row[index_map.get("notes", -1)]) or None if "notes" in index_map else None,
            )
        )

    if not readings:
        raise ReportingError("In der Arbeitsmappe wurden keine reportfähigen Ablesungen gefunden.")
    logger.info("%s reportfähige Ablesungen wurden aus der kombinierten Arbeitsmappe geladen.", len(readings))
    return readings


def _coerce_report_date(value: object) -> date:
    """Wandelt einen Zellenwert aus dem Reporting-Workbook in ein Datum um."""
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    if not text:
        raise ExcelDataError("Im kombinierten Workbook wurde ein leeres reading_date gefunden.")
    return date.fromisoformat(text)


def group_report_readings(readings: Iterable[ReportReading]) -> dict[tuple[str, str], list[ReportReading]]:
    """Gruppiert Ablesungen nach Liegenschaft und Zählerplatz."""
    grouped: dict[tuple[str, str], list[ReportReading]] = defaultdict(list)
    for reading in readings:
        grouped[(reading.property_name, reading.meter_location)].append(reading)
    return dict(grouped)


def segment_readings(
    readings: list[ReportReading],
    max_points_per_band: int,
    min_last_band_points: int,
) -> list[list[ReportReading]]:
    """Segmentiert Ablesungen in Diagrammbänder."""
    if not readings:
        return []
    total = len(readings)
    if total <= max_points_per_band:
        return [list(readings)]

    full_bands = total // max_points_per_band
    remainder = total % max_points_per_band
    band_sizes = [max_points_per_band] * full_bands

    if remainder:
        if remainder < min_last_band_points and band_sizes:
            needed = min_last_band_points - remainder
            donor_index = len(band_sizes) - 1
            while needed > 0 and donor_index >= 0:
                minimum_donor_size = min_last_band_points if len(band_sizes) > 1 else min_last_band_points + 1
                if band_sizes[donor_index] > minimum_donor_size:
                    band_sizes[donor_index] -= 1
                    remainder += 1
                    needed -= 1
                else:
                    donor_index -= 1
        band_sizes.append(remainder)

    segments: list[list[ReportReading]] = []
    cursor = 0
    for size in band_sizes:
        segments.append(readings[cursor : cursor + size])
        cursor += size
    return segments


def split_band_readings(readings: list[ReportReading]) -> tuple[list[ReportReading], list[ReportReading]]:
    """Teilt Band-Ablesungen in linke und rechte Tabellenhälfte."""
    left_count = (len(readings) + 1) // 2
    return readings[:left_count], readings[left_count:]


def enrich_readings_with_metrics(readings: list[ReportReading]) -> None:
    """Ergänzt Tage, Tagesverbrauch und Zählerwechsel-Informationen."""
    previous_reading: ReportReading | None = None
    for reading in readings:
        if previous_reading is None:
            reading.meter_change = False
            reading.days_since_previous = None
            reading.daily_consumption = 0.0
        elif (previous_reading.meter_number or "") != (reading.meter_number or ""):
            reading.meter_change = True
            reading.days_since_previous = None
            reading.daily_consumption = 0.0
        else:
            reading.meter_change = False
            day_delta = (reading.reading_date - previous_reading.reading_date).days
            reading.days_since_previous = day_delta if day_delta > 0 else None
            reading.daily_consumption = (
                (reading.reading_value - previous_reading.reading_value) / day_delta if day_delta > 0 else 0.0
            )
        previous_reading = reading


def paginate_photos(readings: Iterable[ReportReading]) -> list[list[PhotoEntry]]:
    """Paginiert Fotos in Vierergruppen für die Fotoseiten."""
    photos: list[PhotoEntry] = []
    photo_number = 1
    for reading in readings:
        for path in reading.attachment_paths:
            if path.exists():
                photos.append(
                    PhotoEntry(
                        image_path=path,
                        photo_number=photo_number,
                        meter_number=reading.meter_number,
                        reading_date=reading.reading_date,
                    )
                )
                photo_number += 1
    return split_list(photos, 4)


def resolve_cover_image(property_name: str, config: ReportingConfig) -> Path | None:
    """Ermittelt das passende Titelbild für eine Liegenschaft."""
    search_root = config.cover_image_dir
    sanitized = sanitize_filename(property_name)
    candidates = [
        search_root / sanitized / "cover.jpg",
        search_root / sanitized / "cover.jpeg",
        search_root / sanitized / "cover.png",
        search_root / sanitized / "cover.webp",
        search_root / f"{sanitized}.jpg",
        search_root / f"{sanitized}.jpeg",
        search_root / f"{sanitized}.png",
        search_root / f"{sanitized}.webp",
    ]
    for candidate in candidates:
        if candidate.exists():
            logger.info("Für %s wird das liegenschaftsspezifische Titelbild %s verwendet", property_name, candidate)
            return candidate

    property_dir = search_root / sanitized
    if property_dir.exists():
        for extension in ("*.jpg", "*.jpeg", "*.png", "*.webp"):
            matches = sorted(property_dir.glob(extension))
            if matches:
                logger.info("Für %s wird das erste Titelbild aus dem Liegenschaftsordner verwendet: %s", property_name, matches[0])
                return matches[0]

    default_cover = search_root / "default.jpg"
    if default_cover.exists():
        logger.info("Für %s wurde kein spezifisches Titelbild gefunden. Fallback auf %s", property_name, default_cover)
        return default_cover

    if config.placeholder_cover_path and config.placeholder_cover_path.exists():
        logger.warning(
            "Für %s wurde kein spezifisches Titelbild gefunden. Platzhalter %s wird verwendet",
            property_name,
            config.placeholder_cover_path,
        )
        return config.placeholder_cover_path
    logger.warning("Für %s ist kein Titelbild verfügbar. Es wird der interne Fallback im Report verwendet.", property_name)
    return None


def build_meter_reports(readings: list[ReportReading], config: ReportingConfig) -> list[MeterReport]:
    """Erzeugt aus Ablesungen die fachlichen Report-Modelle."""
    reports: list[MeterReport] = []
    generated_on = date.today()
    version_code = generated_on.strftime("%y%m%d")

    grouped = group_report_readings(readings)
    logger.info("%s Berichtsgruppe(n) werden aus den kombinierten Ablesungen aufgebaut.", len(grouped))
    for (property_name, meter_location), grouped_readings in sorted(grouped.items()):
        ordered = sorted(grouped_readings, key=lambda item: (item.reading_date, item.meter_number or "", item.reading_value))
        enrich_readings_with_metrics(ordered)
        segments = segment_readings(
            ordered,
            max_points_per_band=config.max_readings_per_chart_band,
            min_last_band_points=config.min_last_chart_band_points,
        )
        bands = []
        for index, segment in enumerate(segments, start=1):
            left_readings, right_readings = split_band_readings(segment)
            bands.append(
                ReadingBand(
                    band_number=index,
                    readings=segment,
                    left_readings=left_readings,
                    right_readings=right_readings,
                )
            )
        photo_pages = paginate_photos(ordered)
        logger.info(
            "Bericht für %s / %s vorbereitet: %s Ablesungen, %s Diagrammband/-bänder und %s Fotoseite(n).",
            property_name,
            meter_location,
            len(ordered),
            len(bands),
            len(photo_pages),
        )
        reports.append(
            MeterReport(
                property_name=property_name,
                meter_location=meter_location,
                meter_numbers=sorted({item.meter_number for item in ordered if item.meter_number}),
                readings=ordered,
                bands=bands,
                photo_pages=photo_pages,
                cover_image_path=resolve_cover_image(property_name, config),
                generated_on=generated_on,
                version_code=version_code,
            )
        )
    return reports


def build_report_output_path(report: MeterReport, output_dir: Path) -> Path:
    """Baut den vollständigen Ausgabepfad für einen Bericht."""
    property_dir = output_dir / sanitize_filename(report.property_name)
    filename = build_report_filename(report)
    return property_dir / filename


def build_report_filename(report: MeterReport) -> str:
    """Erzeugt den Dateinamen eines Berichts."""
    street_part, locality_part = split_property_filename_parts(report.property_name)
    meter_part = build_meter_location_filename_part(report.meter_location)
    return f"{street_part}__{locality_part}__Zaehlerplatz__{meter_part}.pdf"


def split_property_filename_parts(property_name: str) -> tuple[str, str]:
    """Zerlegt den Liegenschaftsnamen in Straßen- und Ortsanteil."""
    normalized_name = normalize_text(property_name)
    if "," in normalized_name:
        locality_raw, street_raw = [part.strip() for part in normalized_name.split(",", 1)]
        street_part = sanitize_filename_german(street_raw)
        locality_part = sanitize_filename_german(locality_raw)
        if street_part and locality_part:
            return street_part, locality_part
    sanitized = sanitize_filename_german(normalized_name)
    return sanitized, sanitized


def build_meter_location_filename_part(meter_location: str) -> str:
    """Formatiert den Zählerplatz für den Dateinamen."""
    parts = [part for part in normalize_text(meter_location).split("_") if part]
    if not parts:
        return "Unbekannt"
    formatted_parts: list[str] = []
    for part in parts:
        if part.isdigit():
            formatted_parts.append(part)
        else:
            formatted_parts.append(sanitize_filename_german(part.title()))
    return "__".join(formatted_parts)


def sanitize_filename_german(value: str) -> str:
    """Bereinigt einen Dateinamen mit deutscher Umlaut-Transliteration."""
    replacements = str.maketrans(
        {
            "ä": "ae",
            "ö": "oe",
            "ü": "ue",
            "Ä": "Ae",
            "Ö": "Oe",
            "Ü": "Ue",
            "ß": "ss",
        }
    )
    return sanitize_filename(normalize_text(value).translate(replacements))


def run_reporting(config: AppConfig, input_workbook_path: str | None = None) -> list[Path]:
    """Führt den kompletten Reporting-Lauf aus."""
    workbook_path = _resolve_input_workbook_path(config, input_workbook_path)
    logger.info("Report-Erzeugung wird aus der Arbeitsmappe %s gestartet", workbook_path)
    readings = load_combined_readings_workbook(workbook_path)
    reports = build_meter_reports(readings, config.reporting)
    ensure_directory(config.reporting.output_dir)
    logger.info("Berichte werden nach %s geschrieben", config.reporting.output_dir)

    output_paths: list[Path] = []
    for report in reports:
        output_path = build_report_output_path(report, config.reporting.output_dir)
        ensure_directory(output_path.parent)
        logger.info(
            "Bericht für %s / %s wird unter %s erzeugt",
            report.property_name,
            report.meter_location,
            output_path,
        )
        render_report_pdf(report, output_path)
        output_paths.append(output_path)
    logger.info("Die Report-Erzeugung ist abgeschlossen. %s Bericht(e) wurden erzeugt.", len(output_paths))
    return output_paths


def _resolve_input_workbook_path(config: AppConfig, input_override: str | None) -> Path:
    """Bestimmt die Eingabe-Arbeitsmappe für den Reporting-Lauf."""
    if input_override:
        return Path(input_override)
    if config.reporting.input_workbook_path:
        return config.reporting.input_workbook_path
    return config.output_dir / "all_properties__combined_readings.xlsx"


def make_table_rows(report: MeterReport) -> list[dict[str, str]]:
    """Erzeugt die Tabellenzeilen für einen Bericht."""
    rows: list[dict[str, str]] = []
    photo_number = 1
    for reading in report.readings:
        referenced_photos: list[str] = []
        for path in reading.attachment_paths:
            if path.exists():
                referenced_photos.append(str(photo_number))
                photo_number += 1
        rows.append(
            {
                "date": format_date(reading.reading_date),
                "meter_value": format_number(reading.reading_value, 0),
                "days": "Zählerwechsel" if reading.meter_change else (
                    str(reading.days_since_previous) if reading.days_since_previous is not None else "—"
                ),
                "daily_consumption": "Zählerwechsel" if reading.meter_change else format_consumption(reading.daily_consumption or 0.0),
                "meter_number": reading.meter_number or "—",
                "photo_no": ", ".join(referenced_photos) if referenced_photos else "—",
            }
        )
    return rows


def paginate_table_rows(table_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    """Paginiert Tabellenzeilen für ein- oder zweispaltige Seiten."""
    pages: list[dict[str, Any]] = []
    total_rows = len(table_rows)

    if total_rows <= TABLE_ROWS_SINGLE_PAGE_ONE_COL:
        return [{"layout": "one_col", "rows": table_rows}]

    rows_per_page = TABLE_ROWS_PER_COL * TABLE_COLS_PER_PAGE
    for chunk in split_list(table_rows, rows_per_page):
        pages.append(
            {
                "layout": "two_col",
                "left_rows": chunk[:TABLE_ROWS_PER_COL],
                "right_rows": chunk[TABLE_ROWS_PER_COL : TABLE_ROWS_PER_COL * 2],
            }
        )
    return pages


def build_chart_segments(report: MeterReport) -> list[dict[str, Any]]:
    """Erzeugt alle Diagrammsegmente eines Berichts."""
    segments: list[dict[str, Any]] = []
    for band in report.bands:
        plot_data = build_chart_plot_data(band.readings)
        segments.append(
            {
                "index": band.band_number,
                "image_data_uri": build_chart_segment_data_uri(band.readings),
                "date_from": format_date(band.readings[0].reading_date),
                "date_to": format_date(band.readings[-1].reading_date),
                "count": len(band.readings),
                "plot_data": plot_data,
            }
        )
    return segments


def paginate_report_pages(report: MeterReport) -> list[dict[str, Any]]:
    """Baut die Seitenstruktur eines Berichts auf."""
    pages: list[dict[str, Any]] = [{"type": "cover"}]
    table_rows = make_table_rows(report)
    chart_segments = build_chart_segments(report)

    if len(table_rows) <= COMBINED_PAGE_MAX_TABLE_ROWS and len(chart_segments) <= 1:
        pages.append(
            {
                "type": "table_chart_combined",
                "table_rows": table_rows,
                "chart_segments": chart_segments,
            }
        )
    else:
        table_pages = paginate_table_rows(table_rows)
        for index, page in enumerate(table_pages, start=1):
            pages.append(
                {
                    "type": "table",
                    "table_layout": page["layout"],
                    "rows": page.get("rows", []),
                    "left_rows": page.get("left_rows", []),
                    "right_rows": page.get("right_rows", []),
                    "table_page_index": index,
                    "table_page_count": len(table_pages),
                }
            )

        chart_page_chunks = split_list(chart_segments, CHART_ROWS_PER_PAGE)
        for index, chunk in enumerate(chart_page_chunks, start=1):
            pages.append(
                {
                    "type": "charts",
                    "chart_segments": chunk,
                    "chart_page_index": index,
                    "chart_page_count": len(chart_page_chunks),
                }
            )

    for photo_page in report.photo_pages:
        pages.append(
            {
                "type": "photos",
                "photos": [
                    {
                        "photo_number": photo.photo_number,
                        "meter_number": photo.meter_number or "—",
                        "reading_date": format_date(photo.reading_date),
                        "image_data_uri": raster_image_file_to_data_uri(
                            photo.image_path,
                            max_width=2200,
                            max_height=1600,
                            jpeg_quality=88,
                        ),
                    }
                    for photo in photo_page
                ],
            }
        )
    return pages


def build_html_for_report(report: MeterReport) -> str:
    """Rendert das HTML für einen Bericht."""
    pages = paginate_report_pages(report)
    template_text = _load_template_text("report.html.j2")
    css_text = _load_template_text("report.css")
    cover_image_data_uri = (
        raster_image_file_to_data_uri(report.cover_image_path, max_width=2400, max_height=1600, jpeg_quality=90)
        if report.cover_image_path
        else ""
    )
    placeholder_cover_data_uri = _build_placeholder_cover_data_uri()
    report_logo_data_uri = _load_report_logo_data_uri()

    try:
        from jinja2 import BaseLoader, Environment
    except ImportError as exc:  # pragma: no cover - dependency error
        raise ReportingError("Der Reporting-Renderer benötigt das Paket 'jinja2'.") from exc

    environment = Environment(loader=BaseLoader(), autoescape=True)
    template = environment.from_string(template_text)
    property_name = report.property_name
    ref_code = f"BWB-{sanitize_filename(property_name)}"
    meter_numbers_display = ", ".join(report.meter_numbers)
    meter_numbers_summary = build_meter_numbers_summary(report.meter_numbers)

    return template.render(
        styles_css=css_text,
        pages=pages,
        property_name=property_name,
        meter_place=report.meter_location,
        meter_numbers=report.meter_numbers,
        meter_numbers_display=meter_numbers_display,
        meter_numbers_summary=meter_numbers_summary,
        total_entries=len(report.readings),
        period_from=format_date(report.readings[0].reading_date) if report.readings else "",
        period_to=format_date(report.readings[-1].reading_date) if report.readings else "",
        created_date=format_date(report.generated_on),
        version_no=f"1 • {report.version_code}",
        report_title="Wasserzähler-Überprüfung",
        ref_code=ref_code,
        cover_image_data_uri=cover_image_data_uri,
        placeholder_cover_data_uri=placeholder_cover_data_uri,
        report_logo_data_uri=report_logo_data_uri,
    )


def render_pdf_with_playwright(html: str, output_path: Path) -> None:
    """Rendert das Bericht-HTML per Playwright in ein PDF."""
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise ReportingError(
            "Der Reporting-Renderer benötigt 'playwright'. Installieren Sie es mit "
            "'pip install playwright' und führen Sie danach 'python -m playwright install chromium' aus."
        ) from exc

    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page()
            page.set_content(html, wait_until="networkidle")
            page.wait_for_function("window.__summaryFitDone === true")
            page.pdf(
                path=str(output_path),
                print_background=True,
                landscape=True,
                prefer_css_page_size=True,
                scale=1.0,
            )
            browser.close()
    except PlaywrightError as exc:
        raise ReportingError(
            "Playwright konnte Chromium nicht starten. Führen Sie 'python -m playwright install chromium' aus, "
            "um die benötigte Browser-Laufzeit zu installieren."
        ) from exc


def render_report_pdf(report: MeterReport, output_path: Path) -> None:
    """Rendert einen einzelnen Bericht in ein PDF."""
    logger.info(
        "PDF-Bericht für %s / %s wird gerendert",
        report.property_name,
        report.meter_location,
    )
    html = build_html_for_report(report)
    render_pdf_with_playwright(html, output_path)
    logger.info("PDF-Bericht fertig gerendert: %s", output_path)


def build_meter_numbers_summary(meter_numbers: list[str]) -> str:
    """Erzeugt die kompakte Zählernummern-Zusammenfassung."""
    if not meter_numbers:
        return ""
    return ", ".join(meter_numbers)


def build_chart_plot_data(readings: list[ReportReading]) -> dict[str, Any]:
    """Bereitet die Plot-Daten für ein Diagrammsegment auf."""
    x_labels = [format_date(reading.reading_date) for reading in readings]
    y_values = [max(reading.daily_consumption or 0.0, 0.0) for reading in readings]
    meter_numbers = [(reading.meter_number or "Unbekannt").strip() or "Unbekannt" for reading in readings]
    connections: list[dict[str, Any]] = []
    for index in range(len(y_values) - 1):
        if (
            meter_numbers[index] == meter_numbers[index + 1]
            and not readings[index].meter_change
            and not readings[index + 1].meter_change
        ):
            connections.append(
                {
                    "start_index": index,
                    "end_index": index + 1,
                    "start_value": y_values[index],
                    "end_value": y_values[index + 1],
                }
            )

    runs: list[dict[str, Any]] = []
    if meter_numbers:
        run_start = 0
        for index in range(1, len(meter_numbers) + 1):
            if index == len(meter_numbers) or meter_numbers[index] != meter_numbers[run_start]:
                runs.append(
                    {
                        "meter_number": meter_numbers[run_start],
                        "start_index": run_start,
                        "end_index": index - 1,
                    }
                )
                run_start = index

    return {
        "x_labels": x_labels,
        "y_values": y_values,
        "meter_numbers": meter_numbers,
        "connections": connections,
        "runs": runs,
    }


def build_chart_segment_data_uri(readings: list[ReportReading]) -> str:
    """Rendert ein Diagrammsegment als PNG-Data-URI."""
    if not readings:
        raise ValueError("Ein Diagrammsegment darf nicht leer sein.")

    plot_data = build_chart_plot_data(readings)
    x_labels = plot_data["x_labels"]
    y_values = plot_data["y_values"]

    y_peak = max(y_values) if y_values else 1.0
    if y_peak <= 0:
        y_peak = 1.0
    y_min = -max(y_peak * 0.08, 0.12)
    y_max = y_peak * (1 + CHART_Y_PADDING_RATIO + 0.14)

    figure = plt.figure(
        figsize=(CHART_WIDTH_IN, CHART_HEIGHT_PER_ROW_IN),
        dpi=CHART_DPI,
    )
    # Fixed axis rectangles keep the main plot stable while reserving a
    # dedicated lower lane for the meter band and its labels.
    axis = figure.add_axes([0.058, 0.33, 0.93, 0.48])
    band_axis = figure.add_axes([0.058, 0.07, 0.93, 0.12])
    positions = list(range(len(y_values)))

    for connection in plot_data["connections"]:
        axis.plot(
                [positions[connection["start_index"]], positions[connection["end_index"]]],
                [connection["start_value"], connection["end_value"]],
                linewidth=1.2,
                color="black",
                linestyle=(0, (3, 3)),
                zorder=2,
            )

    axis.scatter(positions, y_values, s=28, color="#1f77b4", zorder=3)
    axis.set_ylim(y_min, y_max)
    axis.set_xlim(-0.6, len(positions) - 0.4)
    axis.set_xticks(positions)
    axis.set_xticklabels(x_labels, rotation=0, fontsize=8)
    axis.set_ylabel("m³/Tag", fontsize=9)
    axis.grid(True, axis="y", alpha=0.35, zorder=1)
    axis.set_title(
        f"Tagesverbrauch m³/Tag • {format_date(readings[0].reading_date)} bis {format_date(readings[-1].reading_date)}",
        fontsize=11,
        pad=10,
    )
    axis.tick_params(axis="x", length=3.2, width=0.9, color="#243b53", pad=7)

    for index, value in enumerate(y_values):
        place_below = value >= y_peak * 0.84
        offset = -10 if place_below else (8 if index % 2 == 0 else 13)
        axis.annotate(
            f"{format_consumption(value)} m³/d",
            (positions[index], value),
            textcoords="offset points",
            xytext=(0, offset),
            ha="center",
            va="top" if place_below else "bottom",
            fontsize=7,
            bbox={
                "boxstyle": "round,pad=0.14,rounding_size=0.12",
                "facecolor": "#ffffff",
                "edgecolor": "none",
                "alpha": 0.96,
            },
        )

    band_axis.set_xlim(-0.6, len(positions) - 0.4)
    band_axis.set_ylim(-0.52, 0.34)
    band_axis.set_xticks([])
    band_axis.set_yticks([])
    band_axis.set_facecolor("none")
    band_axis.patch.set_alpha(0)
    axis.set_zorder(2)
    band_axis.set_zorder(1)

    y_band = 0.04
    y_text = -0.24
    for run in plot_data["runs"]:
        start_index = run["start_index"]
        end_index = run["end_index"]
        band_axis.plot(
            [start_index, end_index],
            [y_band, y_band],
            linewidth=1.0,
            color="black",
            zorder=4,
        )
        band_axis.plot(
            [start_index, start_index],
            [y_band - 0.08, y_band + 0.08],
            linewidth=1.0,
            color="black",
            zorder=4,
        )
        band_axis.plot(
            [end_index, end_index],
            [y_band - 0.08, y_band + 0.08],
            linewidth=1.0,
            color="black",
            zorder=4,
        )
        band_axis.text(
            (start_index + end_index) / 2,
            y_text,
            run["meter_number"],
            ha="center",
            va="top",
            fontsize=6.6,
        )

    axis.spines["top"].set_visible(False)
    axis.spines["left"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["bottom"].set_visible(True)
    axis.tick_params(axis="y", length=0)
    band_axis.spines["top"].set_visible(False)
    band_axis.spines["left"].set_visible(False)
    band_axis.spines["right"].set_visible(False)
    band_axis.spines["bottom"].set_visible(False)

    buffer = BytesIO()
    figure.savefig(buffer, format="png", transparent=False)
    plt.close(figure)
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _load_template_text(name: str) -> str:
    """Lädt den Inhalt einer Template-Datei aus dem Paket."""
    template_path = files("watermeter.reporting").joinpath("templates", name)
    return template_path.read_text(encoding="utf-8")


def _load_report_logo_data_uri() -> str:
    """Lädt das Report-Logo als SVG-Data-URI."""
    logo_path = files("watermeter.reporting").joinpath("templates", "immoscan-logo.svg")
    if not logo_path.is_file():
        return ""
    encoded = base64.b64encode(logo_path.read_bytes()).decode("ascii")
    return f"data:image/svg+xml;base64,{encoded}"


def _build_placeholder_cover_data_uri() -> str:
    """Erzeugt ein generiertes Platzhalter-Titelbild als Data-URI."""
    image = Image.new("RGB", (1600, 900), "#d9e2ec")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 900), fill="#d9e2ec")
    draw.rectangle((130, 650, 1470, 820), fill="#bcccdc")
    draw.polygon(
        [(220, 670), (560, 390), (860, 600), (1080, 310), (1380, 670)],
        fill="#a3b7c8",
    )
    draw.rectangle((420, 300, 840, 670), fill="#6d93b8")
    draw.rectangle((920, 190, 1220, 670), fill="#5088B4")
    draw.text((125, 90), "Titelbild nicht verfügbar", fill="#172734")
    return _pil_image_to_data_uri(image)
