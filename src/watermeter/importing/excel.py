from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from ..config import ExcelConfig
from ..exceptions import ExcelDataError
from ..shared.normalization import excel_column_to_index, normalize_meter_location, normalize_text
from .models import MeterReading


logger = logging.getLogger(__name__)


def _coerce_date(value: object) -> date:
    """Wandelt einen Excel-Wert in ein Datum um.

    Args:
        value: Zellwert aus Excel.

    Returns:
        date: Normalisiertes Datum.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ExcelDataError(f"Es wurde ein Datumswert erwartet, erhalten: {value!r}")


def _coerce_float(value: object) -> float:
    """Wandelt einen Excel-Wert in eine numerische Ablesung um.

    Args:
        value: Zellwert aus Excel.

    Returns:
        float: Normalisierte Zahl.
    """
    if isinstance(value, (int, float)):
        return float(value)
    raise ExcelDataError(f"Es wurde ein numerischer Ablesewert erwartet, erhalten: {value!r}")


def _iter_rows(path: Path, config: ExcelConfig) -> Iterable[tuple[int, tuple[object, ...]]]:
    """Iteriert zeilenweise über das konfigurierte Excel-Sheet.

    Args:
        path: Pfad zur Excel-Datei.
        config: Excel-Konfiguration.

    Yields:
        tuple[int, tuple[object, ...]]: Zeilennummer und Zeileninhalt.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    if config.sheet_name and config.sheet_name not in workbook.sheetnames:
        raise ExcelDataError(f"Das konfigurierte Blatt '{config.sheet_name}' existiert nicht.")
    sheet = workbook[config.sheet_name] if config.sheet_name else workbook[workbook.sheetnames[0]]
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        yield row_number, row


def load_historical_readings(path: Path, config: ExcelConfig) -> list[MeterReading]:
    """Liest die historischen Ablesungen aus der Excel-Datei ein.

    Args:
        path: Pfad zur Excel-Datei.
        config: Excel-Konfiguration.

    Returns:
        list[MeterReading]: Geladene historische Ablesungen.
    """
    if not path.exists():
        raise ExcelDataError(f"Die historische Excel-Datei existiert nicht: {path}")

    indexes = {field: excel_column_to_index(column) for field, column in config.columns.items()}
    required = ["meter_number", "reading_date", "reading_value", "property_name", "meter_location"]
    missing = [field for field in required if field not in indexes]
    if missing:
        raise ExcelDataError(f"Fehlende Excel-Spaltenzuordnung(en): {', '.join(missing)}")

    rows: list[MeterReading] = []
    for row_number, row in _iter_rows(path, config):
        if not any(row):
            continue
        try:
            property_name = normalize_text(row[indexes["property_name"]])
            meter_location_raw = normalize_text(row[indexes["meter_location"]])
            reading = MeterReading(
                source="excel_history",
                property_name=property_name,
                meter_location=normalize_meter_location(meter_location_raw),
                raw_meter_location=meter_location_raw,
                meter_number=normalize_text(row[indexes["meter_number"]]) or None,
                reading_date=_coerce_date(row[indexes["reading_date"]]),
                reading_value=_coerce_float(row[indexes["reading_value"]]),
                notes=normalize_text(row[indexes["notes"]]) or None if "notes" in indexes and indexes["notes"] < len(row) else None,
                raw_payload={"row_number": row_number},
            )
        except IndexError as exc:
            raise ExcelDataError(f"Die Zeile {row_number} ist kürzer als die konfigurierte Spaltenzuordnung.") from exc
        except ExcelDataError:
            raise
        except Exception as exc:  # pragma: no cover - defensive
            raise ExcelDataError(f"Zeile {row_number} konnte nicht gelesen werden: {exc}") from exc

        if not property_name or not meter_location_raw:
            logger.warning("Excel-Zeile %s wird übersprungen, weil Liegenschaft oder Zählerplatz leer ist.", row_number)
            continue

        if "consumption_delta" in indexes and indexes["consumption_delta"] < len(row):
            delta = row[indexes["consumption_delta"]]
            if delta is not None:
                reading.extra["consumption_delta"] = delta

        rows.append(reading)

    if not rows:
        raise ExcelDataError("Aus der Excel-Datei konnten keine historischen Ablesungen geladen werden.")
    return rows
